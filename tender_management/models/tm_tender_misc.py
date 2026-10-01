# -*- coding: utf-8 -*-
import base64
import io
from random import randint

from odoo import _, fields, models
from odoo.exceptions import UserError


class TmTenderLostReason(models.Model):
    _name = 'tm.tender.lost.reason'
    _description = 'Tender Lost Reason'
    _order = 'sequence, id'

    name = fields.Char('Reason', required=True, translate=True)
    sequence = fields.Integer(default=10)
    active = fields.Boolean(default=True)


class TmTenderTag(models.Model):
    _name = 'tm.tender.tag'
    _description = 'Tender Tag'

    name = fields.Char('Tag', required=True, translate=True)
    color = fields.Integer('Color', default=lambda self: randint(1, 11))


class TmTenderLostWizard(models.TransientModel):
    _name = 'tm.tender.lost.wizard'
    _description = 'Mark Tender as Lost'

    tender_id = fields.Many2one('tm.tender', required=True)
    reason_id = fields.Many2one('tm.tender.lost.reason', string='Lost Reason', required=True)
    note = fields.Text('Lost Notes')

    def action_confirm(self):
        self.ensure_one()
        self.tender_id._check_manager()
        self.tender_id.write({
            'state': 'lost',
            'lost_reason_id': self.reason_id.id,
            'lost_reason': self.note,
        })
        return {'type': 'ir.actions.act_window_close'}


class TmTenderAddProductsLine(models.TransientModel):
    _name = 'tm.tender.add.products.line'
    _description = 'Add Products Line'

    wizard_id = fields.Many2one('tm.tender.add.products.wizard', required=True, ondelete='cascade')
    product_id = fields.Many2one('product.product', string='Product', required=True)
    quantity = fields.Float('Quantity', default=1.0, digits='Product Unit')


class TmTenderAddProductsWizard(models.TransientModel):
    _name = 'tm.tender.add.products.wizard'
    _description = 'Add Products to Tender'

    tender_id = fields.Many2one('tm.tender', required=True)
    line_ids = fields.One2many('tm.tender.add.products.line', 'wizard_id', string='Products')
    file = fields.Binary('Excel File')
    filename = fields.Char('File Name')

    def _done(self, message=None):
        if message:
            return {
                'type': 'ir.actions.client',
                'tag': 'display_notification',
                'params': {'title': _('Products'), 'message': message, 'type': 'warning',
                           'sticky': True, 'next': {'type': 'ir.actions.act_window_close'}},
            }
        return {'type': 'ir.actions.act_window_close'}

    def _check_tender(self):
        if self.tender_id.state != 'draft':
            raise UserError(_('Products can only be added to a draft tender.'))

    def action_add(self):
        self.ensure_one()
        self._check_tender()
        if not self.line_ids:
            raise UserError(_('Add at least one product.'))
        self.env['tm.tender.line'].create([{
            'tender_id': self.tender_id.id,
            'product_id': line.product_id.id,
            'quantity': line.quantity or 1.0,
        } for line in self.line_ids])
        return self._done()

    def _find_product(self, text):
        Product = self.env['product.product']
        text = str(text).strip()
        return (Product.search([('default_code', '=', text)], limit=1)
                or Product.search([('barcode', '=', text)], limit=1)
                or Product.search([('name', '=', text)], limit=1)
                or Product.search([('name', 'ilike', text)], limit=1))

    def action_import_xlsx(self):
        self.ensure_one()
        self._check_tender()
        if not self.file:
            raise UserError(_('Choose an Excel file first.'))
        try:
            import openpyxl
        except ImportError:
            raise UserError(_('The Python library openpyxl is not installed on the server.'))
        try:
            book = openpyxl.load_workbook(
                io.BytesIO(base64.b64decode(self.file)), data_only=True, read_only=True)
            rows = list(book.active.iter_rows(values_only=True))
        except Exception:  # noqa: BLE001
            raise UserError(_('The file could not be read. Use an .xlsx file.'))
        if rows and rows[0] and isinstance(rows[0][0], str) and rows[0][0].strip().lower() in (
                'product', 'item', 'code'):
            rows = rows[1:]
        vals_list, problems = [], []
        for index, row in enumerate(rows, start=2):
            if not row or row[0] in (None, ''):
                continue
            product = self._find_product(row[0])
            if not product:
                problems.append(_('Row %s: product "%s" not found') % (index, row[0]))
                continue
            try:
                quantity = float(row[1]) if len(row) > 1 and row[1] not in (None, '') else 1.0
            except (TypeError, ValueError):
                problems.append(_('Row %s: invalid quantity') % index)
                continue
            vals = {'tender_id': self.tender_id.id, 'product_id': product.id, 'quantity': quantity}
            if len(row) > 2 and row[2] not in (None, ''):
                try:
                    vals['price_unit'] = float(row[2])
                except (TypeError, ValueError):
                    problems.append(_('Row %s: invalid price') % index)
            vals_list.append(vals)
        if vals_list:
            self.env['tm.tender.line'].create(vals_list)
        if problems:
            shown = '; '.join(problems[:8]) + ('...' if len(problems) > 8 else '')
            return self._done(_('%s products added. Problems: %s') % (len(vals_list), shown))
        if not vals_list:
            raise UserError(_('No products were found in the file.'))
        return self._done()

    def action_download_template(self):
        try:
            import xlsxwriter
        except ImportError:
            raise UserError(_('The Python library xlsxwriter is not installed on the server.'))
        output = io.BytesIO()
        book = xlsxwriter.Workbook(output, {'in_memory': True})
        sheet = book.add_worksheet('Products')
        head = book.add_format({'bold': True, 'bg_color': '#1E3A5F', 'font_color': '#FFFFFF', 'border': 1})
        for col, title in enumerate(('Product', 'Quantity', 'Sale Price (optional)')):
            sheet.write(0, col, title, head)
        sheet.write(1, 0, 'PRODUCT-CODE')
        sheet.write(1, 1, 10)
        sheet.set_column(0, 0, 32)
        sheet.set_column(1, 2, 20)
        book.close()
        attachment = self.env['ir.attachment'].create({
            'name': 'Tender_Products_Template.xlsx',
            'datas': base64.b64encode(output.getvalue()),
            'type': 'binary',
            'mimetype': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'self',
        }
